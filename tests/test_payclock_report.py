import importlib.util
import io
import unittest
import uuid
import shutil
from unittest.mock import patch
from datetime import datetime
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location('report_adapter', ROOT/'payclock-report.py')
adapter = importlib.util.module_from_spec(spec)
spec.loader.exec_module(adapter)
PROJECT = Path('C:/Users/IT MIR/Desktop/payclock')


class ExcelParity(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.engine = adapter.load_engine(PROJECT)

    def assert_excel_parity(self, rows, start, end, office):
        report = adapter.make_report(self.engine, [dict(r) for r in rows], start, end, {}, office)
        selected = [r for r in rows if start <= r['timestamp'][:10] <= end]
        days = self.engine.build_days(selected)
        warehouse, office_days = self.engine.split_days_by_group(days, office)
        for group, group_days in [('Warehouse', warehouse), ('Office', office_days)]:
            # Exercise the actual Excel writer in memory. No source/export files are changed.
            file = io.BytesIO()
            self.engine.write_excel(file, group_days, f'{start} to {end} ({group})')
            file.seek(0)
            book = openpyxl.load_workbook(file, data_only=True)
            for sheet, key, headers in [('Daily Hours', 'daily', 'dailyHeaders'), ('Period Totals', 'totals', 'totalHeaders'), ('Needs Review', 'daily', 'dailyHeaders')]:
                actual = list(book[sheet].values)
                self.assertEqual(list(actual[0]), report[headers])
                expected = [r['values'] for r in report[key] if r['group'] == group and (sheet != 'Needs Review' or not r['values'][15].startswith('ok'))]
                normalize = lambda values: [v.strftime('%Y-%m-%d') if isinstance(v, datetime) else '' if v is None else v for v in values]
                self.assertEqual([normalize(r) for r in actual[1:]], [normalize(r) for r in expected], sheet)
            book.close()
        return report

    def test_real_csv_all_columns_and_totals_match_excel(self):
        rows = self.engine.load_punches(PROJECT/'exports/punches_20260810_to_20260823.csv')
        self.engine.apply_roster_names(rows, self.engine.load_name_map(PROJECT/'employee_name_map.txt'))
        report = self.assert_excel_parity(rows, '2026-08-10', '2026-08-23', self.engine.load_office_ids(PROJECT/'office_employees.txt'))
        self.assertEqual(report['punchCount'], 1716)

    def test_edge_cases_and_inclusive_date_boundaries(self):
        rows = []
        patterns = [
            ('2026-09-01', ['08:00:00','12:00:00','12:30:00','17:15:00']),
            ('2026-09-02', ['08:00:00','16:00:00']),
            ('2026-09-03', ['08:00:00']),
            ('2026-09-04', ['08:00:00','12:00:00','12:30:00']),
            ('2026-09-05', ['08:00:00','10:00:00','11:00:00','13:00:00','14:00:00','15:00:00']),
            ('2026-09-06', ['00:00:00','23:59:59']),
            ('2026-09-07', ['08:00:00','16:00:00']),
        ]
        for date, times in patterns:
            rows.extend({'employee_id':'1','name':'Test','timestamp':date+' '+t} for t in times)
        report = self.assert_excel_parity(rows, '2026-09-01', '2026-09-06', {'1'})
        self.assertEqual(len(report['daily']), 6)
        self.assertIsNone(report['daily'][2]['values'][9])
        self.assertEqual(report['daily'][4]['values'][10], 0)
        self.assertEqual(len(report['daily'][4]['punches']), 6)
        self.assertIn('extra:', report['daily'][4]['values'][16])

    def test_empty_range_has_no_invented_attendance(self):
        report = self.assert_excel_parity([], '2026-09-01', '2026-09-15', set())
        self.assertEqual(report['daily'], [])
        self.assertEqual(report['totals'], [])


class CorrectedReviewExport(unittest.TestCase):
    def test_resolved_correction_stays_in_review_excel_and_is_blue(self):
        engine=adapter.load_engine(PROJECT)
        rows=[{'employee_id':'1','name':'Test employee','timestamp':'2026-10-09 '+t,'source':'local-adjustment','workcode':''} for t in ['08:00:00','12:00:00','13:00:00','17:00:00']]
        report=adapter.make_report(engine,rows,'2026-10-09','2026-10-09',{},set())
        report['daily'][0]['audit']=[{'restoreOriginal':0}]
        self.assertTrue(report['daily'][0]['values'][15].startswith('ok'))
        folder=ROOT/'tests'/'scratch'/('corrected-review-'+uuid.uuid4().hex);folder.mkdir(parents=True)
        try:
            file=folder/'review.xlsx'
            result=adapter.export_excel(engine,{'report':report,'filters':{'sheet':'review','team':'all','employee':'all','day':'','sortColumn':2,'sortDirection':1},'correctedRows':[True],'excelFile':str(file)})
            self.assertEqual(result['exportedRows'],1)
            book=openpyxl.load_workbook(file)
            self.assertEqual(book.active['A2'].fill.fgColor.rgb,'00DCEEFF');book.close()
        finally:shutil.rmtree(folder)


class PersistentCache(unittest.TestCase):
    def setUp(self):
        self.parent = ROOT/'tests'/'scratch'
        self.folder = self.parent/uuid.uuid4().hex
        self.folder.mkdir(parents=True)
        self.engine = adapter.load_engine(PROJECT)
        self.cache = adapter.PunchCache(self.folder/'cache.sqlite3')

    def tearDown(self):
        assert self.folder.resolve().is_relative_to(self.parent.resolve())
        shutil.rmtree(self.folder)

    def punch(self, day, time, employee='1', name='Test'):
        return {'employee_id':employee, 'name':name, 'timestamp':day+' '+time, 'workcode':'', 'source':'face'}

    def test_overlap_replaces_changed_deleted_and_empty_days_and_preserves_outside(self):
        old=[self.punch('2026-09-01','08:00:00'),self.punch('2026-09-01','16:00:00'),self.punch('2026-09-02','08:00:00'),self.punch('2026-09-02','16:00:00'),self.punch('2026-09-03','08:00:00')]
        self.cache.ingest(self.engine,old,'2026-09-01','2026-09-03','old.csv','one')
        new=[self.punch('2026-09-02','09:00:00'),self.punch('2026-09-02','17:00:00')]
        self.cache.ingest(self.engine,new,'2026-09-02','2026-09-04','new.csv','two')
        restarted=adapter.PunchCache(self.folder/'cache.sqlite3')
        rows=restarted.rows('2026-09-01','2026-09-04')
        self.assertEqual([r['timestamp'] for r in rows],['2026-09-01 08:00:00','2026-09-01 16:00:00','2026-09-02 09:00:00','2026-09-02 17:00:00'])
        # Rescanning an old source at startup cannot restore superseded punches.
        self.assertFalse(restarted.ingest(self.engine,old,'2026-09-01','2026-09-03','old.csv','one'))
        self.assertEqual(len(restarted.rows('2026-09-01','2026-09-04')),4)
        self.assertEqual(restarted.info('2026-09-01','2026-09-05')['missingDates'],['2026-09-05'])

    def test_duplicate_records_do_not_double_count_and_bad_import_is_atomic(self):
        punch=self.punch('2026-09-01','08:00:00')
        self.cache.ingest(self.engine,[punch,punch],'2026-09-01','2026-09-01','first','one')
        self.assertEqual(len(self.cache.rows('2026-09-01','2026-09-01')),1)
        with self.assertRaises(ValueError):
            self.cache.ingest(self.engine,[{**punch,'timestamp':'not a timestamp'}],'2026-09-01','2026-09-01','bad','two')
        self.assertEqual(len(self.cache.rows('2026-09-01','2026-09-01')),1)

    def test_late_discovery_of_older_file_cannot_overwrite_newer_dates(self):
        new=[self.punch('2026-09-02','09:00:00')]
        old=[self.punch('2026-09-01','08:00:00'),self.punch('2026-09-02','08:00:00')]
        self.cache.ingest(self.engine,new,'2026-09-02','2026-09-02','new','new',200)
        self.cache.ingest(self.engine,old,'2026-09-01','2026-09-02','old','old',100)
        self.assertEqual([r['timestamp'] for r in self.cache.rows('2026-09-01','2026-09-02')],['2026-09-01 08:00:00','2026-09-02 09:00:00'])

    def test_excel_uses_filtered_snapshot_and_keeps_text_as_text(self):
        rows=[self.punch('2026-09-01','08:00:00',name='=Untrusted'),self.punch('2026-09-01','16:00:00',name='=Untrusted')]
        report=adapter.make_report(self.engine,rows,'2026-09-01','2026-09-02',{},set())
        file=self.folder/'download.xlsx'
        result=adapter.export_excel(self.engine,{'report':report,'filters':{'sheet':'daily','team':'all','employee':'1','day':'2026-09-01','sortColumn':2,'sortDirection':1},'excelFile':str(file)})
        self.assertEqual(result['exportedRows'],1)
        book=openpyxl.load_workbook(file)
        self.assertEqual(book.sheetnames,['Daily Hours'])
        self.assertEqual(book.active['B2'].value,'=Untrusted')
        self.assertEqual(book.active['B2'].data_type,'s')
        self.assertEqual(book.active['J2'].value,8)
        self.assertEqual(book.active.max_column,17)
        book.close()

    def test_automatic_reads_keep_today_partial_until_final_read_tomorrow(self):
        import facein_pull
        request={'project':str(PROJECT),'output':str(self.folder),'source':'auto','host':'unused-test-host',
                 'start':'2026-09-02','end':'2026-09-02','today':'2026-09-02','files':[]}
        commands=[]
        def morning(client,command,timeout=None):
            commands.append(command)
            return 'Return(result="success"\nid="1" name="Test" time="2026-09-02 08:00:00"\nid="1" name="Test" time="2026-09-02 12:00:00")'
        with patch.object(facein_pull.FaceInClient,'send',morning):
            first=adapter.run(request)
        self.assertEqual(first['cache']['partialDates'],['2026-09-02'])
        self.assertEqual(first['cache']['completeDates'],[])
        self.assertEqual(len(commands),1)
        self.assertTrue(commands[0].startswith('GetRecord('))
        # It must remain incomplete when the date changes, even before a read occurs.
        cache=adapter.PunchCache(self.folder/'punches.sqlite3')
        self.assertEqual(cache.info('2026-09-02','2026-09-02','2026-09-03')['missingDates'],['2026-09-02'])
        def evening(client,command,timeout=None):
            commands.append(command)
            return morning(client,command,timeout)[:-1]+'\nid="1" name="Test" time="2026-09-02 12:30:00"\nid="1" name="Test" time="2026-09-02 17:00:00")'
        request['today']='2026-09-03'
        with patch.object(facein_pull.FaceInClient,'send',evening):
            final=adapter.run(request)
        self.assertEqual(final['cache']['completeDates'],['2026-09-02'])
        self.assertEqual(final['punchCount'],4)
        with patch.object(facein_pull.FaceInClient,'send',side_effect=AssertionError('Cached full days must not contact the clock')):
            cached=adapter.run(request)
        self.assertEqual(cached['fetchedRanges'],[])
        self.assertEqual(cached['punchCount'],4)
        self.assertEqual(list(self.folder.glob('*.csv')),[])

    def test_append_only_pull_preserves_existing_timestamps_and_empty_full_days(self):
        old=self.punch('2026-09-01','08:00:00',name='Local name')
        self.cache.ingest(self.engine,[old],'2026-09-01','2026-09-01','first','one',append_only=True)
        self.cache.ingest(self.engine,[{**old,'name':'Changed','source':'different metadata'},self.punch('2026-09-01','17:00:00')],'2026-09-01','2026-09-02','second','two',append_only=True)
        with self.cache.connect() as c:
            self.assertEqual(c.execute('SELECT name FROM punches WHERE timestamp=?',('2026-09-01 08:00:00',)).fetchone()[0],'Local name')
        info=self.cache.info('2026-09-01','2026-09-02')
        self.assertEqual(info['missingDates'],[])
        self.assertIn('2026-09-02',info['completeDates'])
        self.assertEqual(len(self.cache.rows('2026-09-01','2026-09-02')),2)


if __name__ == '__main__':
    unittest.main(verbosity=2)
